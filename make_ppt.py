# --- config --- #
MAX_CHARS = 12  # 1スライドの推奨最大文字数（必要なら変更）
FONT_NAME = "BIZ UDPGothic"  # 例: MS Gothic
SUBTITLE_FILE = "subtitle.txt"
# -------------- #

from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches
import datetime
import os
import sys

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# -------------------------------
# 字幕テキスト読み込み
# 空行で字幕ブロック分割（OS差異対策）
# -------------------------------

if not os.path.exists(SUBTITLE_FILE):
    print(f"エラー: {SUBTITLE_FILE} が見つかりません。")
    print(f"字幕テキストを {SUBTITLE_FILE} として配置してから再実行してください（subtitle_example.txt を参考）。")
    sys.exit(1)

with open(SUBTITLE_FILE, encoding="utf-8") as f:
    lines = f.readlines()

# 空行が連続する場合、区切りの空行に挟まれた 何もない行 は黒スラを出す（空行3つ → 空スライド1枚、5つ → 2枚 …）

blocks = []
current_block = []
blank_count = 0

for line in lines:
    line = line.strip()

    if line == "":
        if current_block:
            blocks.append(current_block)
            current_block = []
        blank_count += 1
    else:
        if blocks:
            blocks.extend([] for _ in range((blank_count - 1) // 2))
        blank_count = 0
        current_block.append(line)

if current_block:
    blocks.append(current_block)

# -------------------------------
# PowerPoint生成
# -------------------------------

for slide_index, subtitle_lines in enumerate(blocks, start=1):

    # 文字数チェック（警告のみ）: スライド全体の文字数で判定
    total_chars = sum(len(line) for line in subtitle_lines)
    if total_chars > MAX_CHARS:
        print(
            f"警告: 字幕{slide_index}枚目が長すぎます "
            f"({total_chars}文字 > {MAX_CHARS}文字) → {' / '.join(subtitle_lines)}"
        )

    # テキストボックスを「スライド上部」に配置
    box_height = int(prs.slide_height * 0.25)
    top = int(prs.slide_height * 0.05)  # 上端から5%の位置に配置

    slide = prs.slides.add_slide(prs.slide_layouts[6])

    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(0, 0, 0)

    # 何もない行のスライドはテキストボックスを入れない
    if not subtitle_lines:
        continue

    box = slide.shapes.add_textbox(
        0,
        top,
        prs.slide_width,
        box_height
    )

    tf = box.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE  # ボックス内でも文字を上下中央に

    for i, line in enumerate(subtitle_lines):

        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()

        p.text = line
        p.font.size = Pt(60)
        p.font.name = FONT_NAME
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = PP_ALIGN.CENTER

# -------------------------------
# 保存
# -------------------------------

nowtime = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
filename = f"subtitles_{nowtime}.pptx"
savepath = f"dist/{filename}"
prs.save(savepath)

print("字幕スライド生成完了:", len(blocks), "枚 -> ", savepath)
